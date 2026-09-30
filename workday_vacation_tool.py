#!/usr/bin/env python3
"""Convert a Workday absence export into Markdown and Outlook calendar events.

The workbook must contain Date, Type, and Status columns. The script detects
their header row, selects approved future vacation, and groups adjacent dates.
Markdown is printed to standard output unless --output is provided.

Requirements
------------
Python 3.10 or newer and openpyxl 3.1.5 or newer.

Install from this repository:

    python -m pip install -r requirements.txt

Alternatively, install the dependency directly when sharing just this script:

    python -m pip install "openpyxl>=3.1.5,<4"

An isolated virtual environment on Windows:

    python -m venv .venv
    .\\.venv\\Scripts\\python -m pip install -r requirements.txt

Usage
-----
Print an aligned Markdown table:

    python workday_vacation_tool.py AbsenceRequests.xlsx

Write Markdown to a file:

    python workday_vacation_tool.py AbsenceRequests.xlsx --output vacation.md

Shell redirection also works:

    python workday_vacation_tool.py AbsenceRequests.xlsx > vacation.md

Override the cutoff date, or explicitly include past dates:

    python workday_vacation_tool.py AbsenceRequests.xlsx --from-date 2030-01-01
    python workday_vacation_tool.py AbsenceRequests.xlsx --include-past

Also create AbsenceRequests.ics for import into Outlook:

    python workday_vacation_tool.py AbsenceRequests.xlsx --ics

Choose a calendar filename:

    python workday_vacation_tool.py AbsenceRequests.xlsx --ics vacation.ics

Use consecutive calendar dates instead of Monday-Friday workday adjacency:

    python workday_vacation_tool.py AbsenceRequests.xlsx --grouping calendar

Run with --help for all options.

Default behavior
----------------
* Includes Type=Vacation and Status=Approved, ignoring case and extra whitespace.
* Includes today and future dates, using the computer's local date.
* Sorts and deduplicates matching dates.
* Groups Friday and the following Monday as one continuous workday range.
* Counts listed vacation dates, not unlisted weekend dates inside a range.
* Does not infer public holidays or support custom workweek schedules.
* Assumes each matching date represents one full vacation day.
* With --ics, writes one calendar containing one all-day event per range.
  Events are titled Planned vacation and carry Outlook out-of-office properties.
  Bridged weekends are covered by these continuous calendar events.

Files are processed locally; nothing is uploaded to Workday or Outlook.
"""

from __future__ import annotations

import argparse
import sys
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from uuid import NAMESPACE_URL, uuid5
from zipfile import BadZipFile

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.utils.exceptions import InvalidFileException
    from openpyxl.worksheet._read_only import ReadOnlyWorksheet
    from openpyxl.worksheet.worksheet import Worksheet
except ModuleNotFoundError as error:
    if error.name != "openpyxl":
        raise
    print(
        'Missing dependency "openpyxl". Install it with:\n'
        '  python -m pip install "openpyxl>=3.1.5,<4"',
        file=sys.stderr,
    )
    raise SystemExit(2) from error


ExcelWorksheet = Worksheet | ReadOnlyWorksheet
REQUIRED_HEADERS = {"date", "type", "status"}
AUTO_ICS_PATH = object()
WEEKDAY_NAMES = ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")
MONTH_NAMES = (
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
)


@dataclass(frozen=True)
class DateRange:
    """An inclusive range and its count of listed vacation dates."""

    start: date
    end: date
    vacation_days: int


@dataclass(frozen=True)
class ConversionResult:
    """Generated content and the files written by a conversion."""

    markdown: str
    markdown_output: Path | None
    ics_output: Path | None


def parse_iso_date(value: str) -> date:
    """Parse a command-line date in YYYY-MM-DD format."""

    try:
        return date.fromisoformat(value)
    except ValueError as error:
        raise argparse.ArgumentTypeError(
            f"{value!r} is not a valid date; expected YYYY-MM-DD"
        ) from error


def build_parser() -> argparse.ArgumentParser:
    """Create the command-line parser."""

    parser = argparse.ArgumentParser(
        description=(
            "Convert an Excel absence export into grouped Markdown vacation ranges."
        ),
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("input", type=Path, help="Path to the source .xlsx workbook")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Write Markdown to this path; omit to print it to standard output",
    )
    parser.add_argument(
        "--sheet",
        help=(
            "Worksheet name. When omitted, the first sheet containing Date, Type, "
            "and Status headers is used"
        ),
    )
    parser.add_argument(
        "--absence-type",
        default="Vacation",
        help="Value to match in the Type column, case-insensitively",
    )
    parser.add_argument(
        "--status",
        default="Approved",
        help="Value to match in the Status column, case-insensitively",
    )

    date_filter = parser.add_mutually_exclusive_group()
    date_filter.add_argument(
        "--from-date",
        type=parse_iso_date,
        help="Earliest date to include in YYYY-MM-DD format; defaults to today",
    )
    date_filter.add_argument(
        "--include-past",
        action="store_true",
        help="Include matching dates before today",
    )

    parser.add_argument(
        "--grouping",
        choices=("workdays", "calendar"),
        default="workdays",
        help=(
            "Use Monday-Friday adjacency or strict calendar-date adjacency when "
            "building ranges"
        ),
    )
    parser.add_argument(
        "--title",
        default="Forthcoming vacation",
        help="Markdown heading written above the table",
    )
    parser.add_argument(
        "--ics",
        dest="ics_output",
        nargs="?",
        const=AUTO_ICS_PATH,
        type=Path,
        metavar="PATH",
        help=(
            "Also create an Outlook-compatible iCalendar file. When PATH is "
            "omitted, use the input filename with an .ics suffix"
        ),
    )
    parser.add_argument(
        "--ics-title",
        default="Planned vacation",
        help="Event title used for every generated iCalendar event",
    )
    return parser


def normalize_text(value: object) -> str:
    """Normalize a workbook value for case-insensitive comparisons."""

    if value is None:
        return ""
    return " ".join(str(value).strip().casefold().split())


def find_header_row(
    worksheet: ExcelWorksheet,
) -> tuple[int, dict[str, int]] | None:
    """Find required headers within the first 50 rows, allowing preamble rows."""

    for row_number, row in enumerate(
        worksheet.iter_rows(max_row=50, values_only=True), start=1
    ):
        columns = {
            normalize_text(value): column_index
            for column_index, value in enumerate(row)
            if value is not None
        }
        if REQUIRED_HEADERS.issubset(columns):
            return row_number, columns
    return None


def select_worksheet(
    workbook: Workbook, sheet_name: str | None
) -> tuple[ExcelWorksheet, int, dict[str, int]]:
    """Select a named worksheet or the first one with the required headers."""

    if sheet_name:
        if sheet_name not in workbook.sheetnames:
            available = ", ".join(workbook.sheetnames)
            raise ValueError(
                f"Worksheet {sheet_name!r} was not found. Available sheets: {available}"
            )
        worksheets = [workbook[sheet_name]]
    else:
        worksheets = workbook.worksheets

    for worksheet in worksheets:
        header = find_header_row(worksheet)
        if header:
            header_row, columns = header
            return worksheet, header_row, columns

    scope = f"worksheet {sheet_name!r}" if sheet_name else "the workbook"
    headers = ", ".join(sorted(name.title() for name in REQUIRED_HEADERS))
    raise ValueError(
        f"Could not find the required headers ({headers}) in the first 50 rows of "
        f"{scope}."
    )


def value_at(row: Sequence[object], column_index: int) -> object:
    """Return a value from a sparse row; missing cells are validated by callers."""

    return row[column_index] if column_index < len(row) else None


def workbook_date(value: object, row_number: int) -> date:
    """Convert an Excel date cell or ISO-formatted text cell to a date."""

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        try:
            return date.fromisoformat(value.strip())
        except ValueError:
            pass

    raise ValueError(
        f"Row {row_number} has an unsupported Date value {value!r}. "
        "Use an Excel date cell or YYYY-MM-DD text."
    )


def extract_vacation_dates(
    worksheet: ExcelWorksheet,
    header_row: int,
    columns: dict[str, int],
    absence_type: str,
    status: str,
    earliest_date: date,
) -> list[date]:
    """Filter, sort, and deduplicate matching dates from the worksheet."""

    expected_type = normalize_text(absence_type)
    expected_status = normalize_text(status)
    matching_dates: set[date] = set()

    rows = worksheet.iter_rows(min_row=header_row + 1, values_only=True)
    for row_number, row in enumerate(rows, start=header_row + 1):
        raw_type = value_at(row, columns["type"])
        raw_status = value_at(row, columns["status"])

        if normalize_text(raw_type) != expected_type:
            continue
        if normalize_text(raw_status) != expected_status:
            continue

        raw_date = value_at(row, columns["date"])
        if raw_date is None:
            raise ValueError(
                f"Row {row_number} matches the filters but has an empty Date cell."
            )

        vacation_date = workbook_date(raw_date, row_number)
        if vacation_date >= earliest_date:
            matching_dates.add(vacation_date)

    return sorted(matching_dates)


def next_workday(day: date) -> date:
    """Return the next Monday-Friday date after day."""

    candidate = day + timedelta(days=1)
    while candidate.weekday() >= 5:
        candidate += timedelta(days=1)
    return candidate


def group_dates(dates: Sequence[date], grouping: str) -> list[DateRange]:
    """Group sorted, unique dates using workday or calendar adjacency."""

    ranges: list[DateRange] = []
    for current_date in dates:
        if ranges:
            previous = ranges[-1]
            expected_date = (
                next_workday(previous.end)
                if grouping == "workdays"
                else previous.end + timedelta(days=1)
            )
            if current_date == expected_date:
                ranges[-1] = DateRange(
                    start=previous.start,
                    end=current_date,
                    vacation_days=previous.vacation_days + 1,
                )
                continue

        ranges.append(DateRange(start=current_date, end=current_date, vacation_days=1))

    return ranges


def format_day(day: date) -> str:
    """Format a date with English names independent of the system locale."""

    return (
        f"{WEEKDAY_NAMES[day.weekday()]} {day.day} "
        f"{MONTH_NAMES[day.month - 1]} {day.year}"
    )


def format_date_range(date_range: DateRange) -> str:
    """Format a single date or a compact inclusive range."""

    start = date_range.start
    end = date_range.end
    if start == end:
        return format_day(start)

    start_weekday = WEEKDAY_NAMES[start.weekday()]
    end_weekday = WEEKDAY_NAMES[end.weekday()]
    start_month = MONTH_NAMES[start.month - 1]
    end_month = MONTH_NAMES[end.month - 1]

    if start.year == end.year and start.month == end.month:
        return (
            f"{start_weekday} {start.day} - {end_weekday} {end.day} "
            f"{end_month} {end.year}"
        )
    if start.year == end.year:
        return (
            f"{start_weekday} {start.day} {start_month} - "
            f"{end_weekday} {end.day} {end_month} {end.year}"
        )
    return (
        f"{start_weekday} {start.day} {start_month} {start.year} - "
        f"{end_weekday} {end.day} {end_month} {end.year}"
    )


def render_markdown(ranges: Sequence[DateRange], title: str, grouping: str) -> str:
    """Render a valid Markdown table with aligned raw-text columns."""

    total_days = sum(date_range.vacation_days for date_range in ranges)
    table_rows = [
        (format_date_range(date_range), str(date_range.vacation_days))
        for date_range in ranges
    ]
    table_rows.append(("**Total**", f"**{total_days}**"))

    date_header = "Date range"
    days_header = "Vacation days"
    date_width = max(len(date_header), *(len(label) for label, _ in table_rows))
    days_width = max(len(days_header), *(len(count) for _, count in table_rows))

    lines = [
        f"# {title}",
        "",
        f"| {date_header.ljust(date_width)} | {days_header.rjust(days_width)} |",
        f"| {'-' * date_width} | {('-' * (days_width - 1)) + ':'} |",
    ]
    lines.extend(
        f"| {label.ljust(date_width)} | {count.rjust(days_width)} |"
        for label, count in table_rows
    )
    lines.append("")

    if grouping == "workdays":
        lines.append(
            "Consecutive workdays are grouped together. Weekend days within a "
            "range are not included in the vacation-day count. Public holidays "
            "are not inferred."
        )
    else:
        lines.append(
            "Only consecutive calendar dates are grouped together. The count is "
            "the number of listed vacation dates."
        )

    return "\n".join(lines) + "\n"


def escape_ics_text(value: str) -> str:
    """Escape text according to the RFC 5545 iCalendar text-value rules."""

    return (
        value.replace("\\", "\\\\")
        .replace("\r\n", "\\n")
        .replace("\r", "\\n")
        .replace("\n", "\\n")
        .replace(";", "\\;")
        .replace(",", "\\,")
    )


def fold_ics_line(line: str, limit: int = 75) -> str:
    """Fold at 75 UTF-8 octets, preserving complete characters and CRLF."""

    segments: list[str] = []
    current = ""

    for character in line:
        candidate = current + character
        if current and len(candidate.encode("utf-8")) > limit:
            segments.append(current)
            current = " " + character
        else:
            current = candidate

    segments.append(current)
    return "\r\n".join(segments)


def render_ics(
    ranges: Sequence[DateRange],
    event_title: str,
    source_identity: str,
) -> str:
    """Render one calendar with one all-day Outlook OOF event per range."""

    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    escaped_title = escape_ics_text(event_title)
    lines = [
        "BEGIN:VCALENDAR",
        "PRODID:-//Workday Vacation Tool//EN",
        "VERSION:2.0",
        "CALSCALE:GREGORIAN",
        "METHOD:PUBLISH",
        f"X-WR-CALNAME:{escaped_title}",
    ]

    for date_range in ranges:
        uid_source = (
            f"{source_identity}|{date_range.start.isoformat()}|"
            f"{date_range.end.isoformat()}|{event_title}"
        )
        event_uid = f"{uuid5(NAMESPACE_URL, uid_source)}@vacation.local"
        exclusive_end = date_range.end + timedelta(days=1)

        lines.extend(
            [
                "BEGIN:VEVENT",
                f"UID:{event_uid}",
                f"DTSTAMP:{timestamp}",
                f"DTSTART;VALUE=DATE:{date_range.start:%Y%m%d}",
                f"DTEND;VALUE=DATE:{exclusive_end:%Y%m%d}",
                f"SUMMARY:{escaped_title}",
                "CATEGORIES:Vacation",
                "CLASS:PUBLIC",
                "STATUS:CONFIRMED",
                "SEQUENCE:0",
                "TRANSP:OPAQUE",
                "X-MICROSOFT-CDO-BUSYSTATUS:OOF",
                "X-MICROSOFT-CDO-INTENDEDSTATUS:OOF",
                "X-MICROSOFT-CDO-ALLDAYEVENT:TRUE",
                "END:VEVENT",
            ]
        )

    lines.append("END:VCALENDAR")
    return "\r\n".join(fold_ics_line(line) for line in lines) + "\r\n"


def convert(args: argparse.Namespace) -> ConversionResult:
    """Read the workbook and write any explicitly requested output files."""

    input_path = args.input.expanduser()
    output_path = args.output.expanduser() if args.output else None
    if args.ics_output is AUTO_ICS_PATH:
        ics_output_path = input_path.with_suffix(".ics")
    elif args.ics_output is not None:
        ics_output_path = args.ics_output.expanduser()
    else:
        ics_output_path = None

    if not input_path.is_file():
        raise FileNotFoundError(f"Input workbook was not found: {input_path}")
    if input_path.suffix.casefold() != ".xlsx":
        raise ValueError(f"Input must be an .xlsx workbook: {input_path}")
    if output_path and input_path.resolve() == output_path.resolve():
        raise ValueError("The output path must be different from the input workbook.")
    if ics_output_path and input_path.resolve() == ics_output_path.resolve():
        raise ValueError("The iCalendar path must differ from the input workbook.")
    if (
        output_path
        and ics_output_path
        and output_path.resolve() == ics_output_path.resolve()
    ):
        raise ValueError("The Markdown and iCalendar output paths must differ.")

    earliest_date = date.min if args.include_past else args.from_date or date.today()

    workbook = openpyxl.load_workbook(input_path, data_only=True, read_only=True)
    try:
        worksheet, header_row, columns = select_worksheet(workbook, args.sheet)
        source_identity = f"{input_path.name}|{worksheet.title}"
        dates = extract_vacation_dates(
            worksheet=worksheet,
            header_row=header_row,
            columns=columns,
            absence_type=args.absence_type,
            status=args.status,
            earliest_date=earliest_date,
        )
    finally:
        workbook.close()

    if not dates:
        raise ValueError(
            "No matching vacation dates were found. Check --from-date, "
            "--absence-type, --status, and --sheet."
        )

    ranges = group_dates(dates, args.grouping)
    markdown = render_markdown(ranges, args.title, args.grouping)
    if output_path:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(markdown, encoding="utf-8")

    if ics_output_path:
        calendar = render_ics(ranges, args.ics_title, source_identity)
        ics_output_path.parent.mkdir(parents=True, exist_ok=True)
        with ics_output_path.open("w", encoding="utf-8", newline="") as ics_file:
            ics_file.write(calendar)

    return ConversionResult(
        markdown=markdown,
        markdown_output=output_path,
        ics_output=ics_output_path,
    )


def main(argv: Sequence[str] | None = None) -> int:
    """Run the CLI; keep errors and file notifications off Markdown stdout."""

    parser = build_parser()
    args = parser.parse_args(argv)

    try:
        result = convert(args)
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        print(f"Error: {error}", file=sys.stderr)
        return 1

    if result.markdown_output:
        print(
            f"Wrote Markdown vacation table to: {result.markdown_output}",
            file=sys.stderr,
        )
    else:
        sys.stdout.write(result.markdown)
    if result.ics_output:
        print(f"Wrote Outlook calendar to: {result.ics_output}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
