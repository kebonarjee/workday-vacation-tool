"""Exercise the shareable CLI using synthetic, nonpersonal absence data."""

from __future__ import annotations

import argparse
import subprocess
import sys
from collections.abc import Sequence
from datetime import date, datetime
from pathlib import Path

import pytest
from icalendar import Calendar
from openpyxl import Workbook

import workday_vacation_tool as tool

SCRIPT = Path(tool.__file__).resolve()


def make_workbook(
    tmp_path: Path,
    rows: Sequence[Sequence[object]],
    headers: Sequence[str] = ("Date", "Type", "Status"),
) -> Path:
    """Write a small export with a Workday-style title before its headers."""

    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Absences"
    worksheet.append(["Absence Requests"])
    worksheet.append(list(headers))
    for row in rows:
        worksheet.append(list(row))
    path = tmp_path / "AbsenceRequests.xlsx"
    workbook.save(path)
    workbook.close()
    return path


def run_cli(workbook: Path, *options: str) -> subprocess.CompletedProcess[str]:
    """Run the actual script from outside the repository, as a colleague would."""

    return subprocess.run(
        [sys.executable, "-B", str(SCRIPT), str(workbook), *options],
        cwd=workbook.parent,
        capture_output=True,
        text=True,
        encoding="utf-8",
        check=False,
    )


def parse_args(workbook: Path, *options: str) -> argparse.Namespace:
    return tool.build_parser().parse_args([str(workbook), *options])


@pytest.mark.parametrize(
    ("value", "expected"),
    [(None, ""), (" Approved ", "approved"), ("  Time   Off ", "time off")],
)
def test_normalize_text(value: object, expected: str) -> None:
    assert tool.normalize_text(value) == expected


def test_parse_iso_date() -> None:
    assert tool.parse_iso_date("2030-05-06") == date(2030, 5, 6)


@pytest.mark.parametrize("value", ["not-a-date", "2030-02-30", "06/05/2030"])
def test_invalid_cli_date(value: str) -> None:
    with pytest.raises(argparse.ArgumentTypeError, match="YYYY-MM-DD"):
        tool.parse_iso_date(value)


@pytest.mark.parametrize(
    "value",
    [date(2030, 5, 6), datetime(2030, 5, 6, 12, 30), " 2030-05-06 "],
)
def test_supported_workbook_dates(value: object) -> None:
    assert tool.workbook_date(value, 7) == date(2030, 5, 6)


@pytest.mark.parametrize("value", [None, 45000, "06/05/2030", "invalid"])
def test_unsupported_workbook_date_has_row_number(value: object) -> None:
    with pytest.raises(ValueError, match="Row 7"):
        tool.workbook_date(value, 7)


def test_select_worksheet_detects_reordered_case_insensitive_headers() -> None:
    workbook = Workbook()
    cover = workbook.active
    assert cover is not None
    cover.title = "Cover"
    cover.append(["Not an absence list"])
    worksheet = workbook.create_sheet("Data")
    worksheet.append(["Export title"])
    worksheet.append([" Status ", "DATE", " type ", "Extra column"])

    selected, row_number, columns = tool.select_worksheet(workbook, None)
    assert selected is worksheet
    assert row_number == 2
    assert columns["date"] == 1
    assert columns["type"] == 2
    assert columns["status"] == 0
    workbook.close()


def test_requested_sheet_and_missing_sheet() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = "Chosen"
    worksheet.append(["Date", "Type", "Status"])
    assert tool.select_worksheet(workbook, "Chosen")[0] is worksheet
    with pytest.raises(ValueError, match="Available sheets: Chosen"):
        tool.select_worksheet(workbook, "Missing")
    workbook.close()


def test_missing_headers_are_reported() -> None:
    workbook = Workbook()
    with pytest.raises(ValueError, match="required headers"):
        tool.select_worksheet(workbook, None)
    workbook.close()


def test_extract_filters_sorts_and_deduplicates_dates() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Date", "Type", "Status"])
    for row in [
        (date(2030, 5, 8), "Vacation", "Approved"),
        (date(2030, 5, 6), " VACATION ", " approved "),
        (date(2030, 5, 6), "Vacation", "Approved"),
        (date(2030, 5, 5), "Vacation", "Approved"),
        (date(2030, 5, 7), "Vacation", "Submitted"),
        (date(2030, 5, 7), "Sick Leave", "Approved"),
        (None, None, None),
    ]:
        worksheet.append(row)

    dates = tool.extract_vacation_dates(
        worksheet,
        1,
        {"date": 0, "type": 1, "status": 2},
        "Vacation",
        "Approved",
        date(2030, 5, 6),
    )
    assert dates == [date(2030, 5, 6), date(2030, 5, 8)]
    workbook.close()


def test_missing_date_on_matching_row_is_reported() -> None:
    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.append(["Date", "Type", "Status"])
    worksheet.append([None, "Vacation", "Approved"])
    with pytest.raises(ValueError, match="Row 2.*empty Date"):
        tool.extract_vacation_dates(
            worksheet,
            1,
            {"date": 0, "type": 1, "status": 2},
            "Vacation",
            "Approved",
            date.min,
        )
    workbook.close()


def test_default_cutoff_ignores_past_and_includes_today(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    class FixedDate(date):
        @classmethod
        def today(cls) -> date:
            return cls(2030, 5, 6)

    workbook = make_workbook(
        tmp_path,
        [
            (date(2030, 5, 5), "Vacation", "Approved"),
            (date(2030, 5, 6), "Vacation", "Approved"),
            (date(2030, 5, 7), "Vacation", "Approved"),
        ],
    )
    monkeypatch.setattr(tool, "date", FixedDate)
    result = tool.convert(parse_args(workbook))
    assert "Mon 6 - Tue 7 May 2030" in result.markdown
    assert "**2**" in result.markdown
    assert "Sun 5" not in result.markdown


@pytest.mark.parametrize(
    ("options", "total"),
    [
        (("--from-date", "2030-05-07"), 1),
        (("--include-past",), 2),
    ],
)
def test_explicit_cutoff_overrides(
    tmp_path: Path, options: tuple[str, ...], total: int
) -> None:
    workbook = make_workbook(
        tmp_path,
        [
            (date(2030, 5, 6), "Vacation", "Approved"),
            (date(2030, 5, 7), "Vacation", "Approved"),
        ],
    )
    result = tool.convert(parse_args(workbook, *options))
    assert f"**{total}**" in result.markdown


def test_date_filter_options_are_mutually_exclusive(tmp_path: Path) -> None:
    with pytest.raises(SystemExit) as error:
        parse_args(
            tmp_path / "export.xlsx", "--include-past", "--from-date", "2030-01-01"
        )
    assert error.value.code == 2


def test_workday_grouping_bridges_weekends_without_counting_them() -> None:
    dates = [date(2030, 5, 10), date(2030, 5, 13), date(2030, 5, 14)]
    assert tool.group_dates(dates, "workdays") == [
        tool.DateRange(date(2030, 5, 10), date(2030, 5, 14), 3)
    ]


def test_calendar_grouping_does_not_bridge_weekends() -> None:
    dates = [date(2030, 5, 10), date(2030, 5, 13), date(2030, 5, 14)]
    assert tool.group_dates(dates, "calendar") == [
        tool.DateRange(date(2030, 5, 10), date(2030, 5, 10), 1),
        tool.DateRange(date(2030, 5, 13), date(2030, 5, 14), 2),
    ]


def test_missing_weekday_splits_ranges() -> None:
    dates = [date(2030, 5, 6), date(2030, 5, 8)]
    assert len(tool.group_dates(dates, "workdays")) == 2
    assert tool.group_dates([], "workdays") == []


@pytest.mark.parametrize(
    ("start", "end", "expected"),
    [
        (date(2030, 5, 6), date(2030, 5, 6), "Mon 6 May 2030"),
        (date(2030, 5, 6), date(2030, 5, 8), "Mon 6 - Wed 8 May 2030"),
        (date(2030, 5, 31), date(2030, 6, 3), "Fri 31 May - Mon 3 June 2030"),
        (
            date(2030, 12, 31),
            date(2031, 1, 1),
            "Tue 31 December 2030 - Wed 1 January 2031",
        ),
    ],
)
def test_format_date_range(start: date, end: date, expected: str) -> None:
    assert tool.format_date_range(tool.DateRange(start, end, 1)) == expected


@pytest.mark.parametrize("count", [3, 123456789012345])
def test_markdown_pipes_align_and_count_column_is_right_aligned(count: int) -> None:
    ranges = [tool.DateRange(date(2030, 5, 6), date(2030, 5, 8), count)]
    markdown = tool.render_markdown(ranges, "Vacation", "workdays")
    table = [line for line in markdown.splitlines() if line.startswith("|")]
    pipe_positions = [
        [index for index, character in enumerate(line) if character == "|"]
        for line in table
    ]
    assert all(positions == pipe_positions[0] for positions in pipe_positions)
    assert table[1].split("|")[2].strip().endswith(":")
    assert table[2].split("|")[2].endswith(f"{count} ")
    assert table[-1].split("|")[1].strip() == "**Total**"
    assert table[-1].split("|")[2].strip() == f"**{count}**"


def test_calendar_grouping_note() -> None:
    markdown = tool.render_markdown([], "Vacation", "calendar")
    assert "consecutive calendar dates" in markdown
    assert "**0**" in markdown


def test_ics_parses_multiple_all_day_out_of_office_events() -> None:
    ranges = [
        tool.DateRange(date(2030, 5, 6), date(2030, 5, 8), 3),
        tool.DateRange(date(2030, 5, 10), date(2030, 5, 13), 2),
    ]
    content = tool.render_ics(ranges, "Planned vacation", "synthetic-export")
    calendar = Calendar.from_ical(content)
    events = calendar.walk("VEVENT")
    assert len(events) == 2
    assert str(calendar["VERSION"]) == "2.0"
    assert events[0].decoded("DTSTART") == date(2030, 5, 6)
    assert events[0].decoded("DTEND") == date(2030, 5, 9)
    assert events[1].decoded("DTSTART") == date(2030, 5, 10)
    assert events[1].decoded("DTEND") == date(2030, 5, 14)
    for event in events:
        assert str(event["SUMMARY"]) == "Planned vacation"
        assert str(event["TRANSP"]) == "OPAQUE"
        assert str(event["X-MICROSOFT-CDO-BUSYSTATUS"]) == "OOF"
        assert str(event["X-MICROSOFT-CDO-INTENDEDSTATUS"]) == "OOF"
        assert str(event["X-MICROSOFT-CDO-ALLDAYEVENT"]) == "TRUE"
        assert event.decoded("DTSTAMP").utcoffset().total_seconds() == 0
    assert len({str(event["UID"]) for event in events}) == 2

    repeated = Calendar.from_ical(
        tool.render_ics(ranges, "Planned vacation", "synthetic-export")
    )
    assert [str(event["UID"]) for event in repeated.walk("VEVENT")] == [
        str(event["UID"]) for event in events
    ]


def test_ics_single_day_has_exclusive_next_day_end() -> None:
    content = tool.render_ics(
        [tool.DateRange(date(2030, 5, 6), date(2030, 5, 6), 1)],
        "Planned vacation",
        "synthetic-export",
    )
    event = Calendar.from_ical(content).walk("VEVENT")[0]
    assert event.decoded("DTEND") == date(2030, 5, 7)


def test_ics_escapes_and_folds_unicode_title_without_corruption() -> None:
    title = "Vacation, rest; travel\\home\n" + "\u00f8" * 80
    content = tool.render_ics(
        [tool.DateRange(date(2030, 5, 6), date(2030, 5, 6), 1)],
        title,
        "synthetic-export",
    )
    encoded = content.encode("utf-8")
    assert b"\r\r\n" not in encoded
    assert encoded.endswith(b"END:VCALENDAR\r\n")
    assert b"\n" not in encoded.replace(b"\r\n", b"")
    assert all(len(line) <= 75 for line in encoded.split(b"\r\n"))
    for line in encoded.split(b"\r\n"):
        line.decode("utf-8")
    unfolded = content.replace("\r\n ", "")
    assert f"X-WR-CALNAME:{tool.escape_ics_text(title)}\r\n" in unfolded
    calendar = Calendar.from_ical(encoded)
    assert str(calendar.walk("VEVENT")[0]["SUMMARY"]) == title


def test_cli_stdout_does_not_create_a_markdown_file(tmp_path: Path) -> None:
    workbook = make_workbook(tmp_path, [(date(2030, 5, 6), "Vacation", "Approved")])
    result = run_cli(workbook, "--include-past")
    assert result.returncode == 0, result.stderr
    assert result.stdout.startswith("# Forthcoming vacation\n\n")
    assert "Mon 6 May 2030" in result.stdout
    assert result.stderr == ""
    assert not workbook.with_suffix(".md").exists()
    assert not workbook.with_suffix(".ics").exists()


def test_cli_explicit_markdown_output(tmp_path: Path) -> None:
    workbook = make_workbook(tmp_path, [(date(2030, 5, 6), "Vacation", "Approved")])
    output = tmp_path / "output" / "vacation.md"
    result = run_cli(workbook, "--include-past", "--output", str(output))
    assert result.returncode == 0, result.stderr
    assert result.stdout == ""
    assert "Wrote Markdown vacation table" in result.stderr
    assert output.read_text(encoding="utf-8").startswith("# Forthcoming vacation")


@pytest.mark.parametrize("explicit_path", [False, True])
def test_cli_calendar_keeps_markdown_stdout_clean(
    tmp_path: Path, explicit_path: bool
) -> None:
    workbook = make_workbook(tmp_path, [(date(2030, 5, 6), "Vacation", "Approved")])
    output = (
        tmp_path / "output" / "vacation.ics"
        if explicit_path
        else workbook.with_suffix(".ics")
    )
    options = ["--include-past", "--ics"]
    if explicit_path:
        options.append(str(output))
    result = run_cli(workbook, *options)
    assert result.returncode == 0, result.stderr
    assert result.stdout.startswith("# Forthcoming vacation\n\n")
    assert "Wrote Outlook calendar" not in result.stdout
    assert "Wrote Outlook calendar" in result.stderr
    encoded = output.read_bytes()
    assert b"\r\r\n" not in encoded
    assert b"\n" not in encoded.replace(b"\r\n", b"")
    assert len(Calendar.from_ical(encoded).walk("VEVENT")) == 1


def test_cli_custom_filters_titles_and_sheet(tmp_path: Path) -> None:
    workbook = make_workbook(
        tmp_path, [(date(2030, 5, 6), "Other absence", "Submitted")]
    )
    calendar = tmp_path / "custom.ics"
    result = run_cli(
        workbook,
        "--include-past",
        "--sheet",
        "Absences",
        "--absence-type",
        "Other absence",
        "--status",
        "Submitted",
        "--title",
        "Time off",
        "--ics",
        str(calendar),
        "--ics-title",
        "Custom vacation",
    )
    assert result.returncode == 0, result.stderr
    assert result.stdout.startswith("# Time off\n")
    event = Calendar.from_ical(calendar.read_bytes()).walk("VEVENT")[0]
    assert str(event["SUMMARY"]) == "Custom vacation"


@pytest.mark.parametrize("option", ["--output", "--ics"])
def test_cli_refuses_to_overwrite_input(tmp_path: Path, option: str) -> None:
    workbook = make_workbook(tmp_path, [(date(2030, 5, 6), "Vacation", "Approved")])
    original = workbook.read_bytes()
    result = run_cli(workbook, "--include-past", option, str(workbook))
    assert result.returncode == 1
    assert result.stdout == ""
    assert "Error:" in result.stderr
    assert workbook.read_bytes() == original


def test_cli_refuses_colliding_output_paths(tmp_path: Path) -> None:
    workbook = make_workbook(tmp_path, [(date(2030, 5, 6), "Vacation", "Approved")])
    output = tmp_path / "same-output"
    result = run_cli(
        workbook, "--include-past", "--output", str(output), "--ics", str(output)
    )
    assert result.returncode == 1
    assert "output paths must differ" in result.stderr
    assert not output.exists()


def test_cli_reports_no_matching_dates_without_writing_outputs(tmp_path: Path) -> None:
    workbook = make_workbook(tmp_path, [(date(2030, 5, 6), "Vacation", "Submitted")])
    output = tmp_path / "vacation.md"
    result = run_cli(workbook, "--include-past", "--output", str(output), "--ics")
    assert result.returncode == 1
    assert "No matching vacation dates" in result.stderr
    assert result.stdout == ""
    assert not output.exists()
    assert not workbook.with_suffix(".ics").exists()


def test_cli_reports_missing_input(tmp_path: Path) -> None:
    result = run_cli(tmp_path / "missing.xlsx")
    assert result.returncode == 1
    assert "Input workbook was not found" in result.stderr
    assert result.stdout == ""


def test_cli_reports_corrupt_workbook(tmp_path: Path) -> None:
    workbook = tmp_path / "corrupt.xlsx"
    workbook.write_text("not an Excel workbook", encoding="utf-8")
    result = run_cli(workbook)
    assert result.returncode == 1
    assert "Error:" in result.stderr
    assert "Traceback" not in result.stderr


def test_cli_help_documents_stdout_and_calendar_options() -> None:
    result = subprocess.run(
        [sys.executable, "-B", str(SCRIPT), "--help"],
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0
    assert "standard output" in result.stdout
    assert "--ics [PATH]" in result.stdout
    assert "--ics-title" in result.stdout
